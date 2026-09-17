from django.contrib import admin
from django.http import HttpResponse
from django.urls import path, reverse
from django.utils.html import format_html

from .models import (
    PublicationCategory, Publication, PublicationImage, PublicationVideo,
    EventRegistration,
)


# ─────────────────────────────
# Images / vidéos inline
# ─────────────────────────────
class PublicationImageInline(admin.TabularInline):
    model  = PublicationImage
    extra  = 3
    fields = ['image', 'caption', 'order']


class PublicationVideoInline(admin.TabularInline):
    model  = PublicationVideo
    extra  = 1
    fields = ['video', 'caption', 'order']


# ─────────────────────────────
# Catégorie
# ─────────────────────────────
@admin.register(PublicationCategory)
class PublicationCategoryAdmin(admin.ModelAdmin):
    list_display        = ['name', 'icon', 'order']
    prepopulated_fields = {'slug': ('name',)}
    ordering            = ['order', 'name']


# ─────────────────────────────
# 📤 Export Excel des inscriptions
# ─────────────────────────────
EXPORT_HEADERS = ['Événement', 'Prénom', 'Nom', 'Email', 'Indicatif', 'Numéro WhatsApp', 'Pays', 'Ville', 'Date d\'inscription']


def _build_registrations_workbook(queryset, sheet_title="Inscriptions"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # limite Excel

    ws.append(EXPORT_HEADERS)
    header_fill = PatternFill(start_color="0D3B7A", end_color="0D3B7A", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for reg in queryset.select_related('publication'):
        ws.append([
            reg.publication.title, reg.first_name, reg.last_name, reg.email,
            reg.phone_code, reg.phone_number, reg.country, reg.city,
            reg.created_at.strftime('%d/%m/%Y %H:%M'),
        ])

    for col in ws.columns:
        width = max(len(str(c.value)) if c.value else 0 for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max(width, 12), 45)

    return wb


def _xlsx_response(wb, filename):
    # Les en-têtes HTTP n'acceptent que de l'ASCII : on fournit un nom de
    # secours sans accents pour les vieux navigateurs, et le vrai nom
    # (accents compris) via le paramètre filename* normalisé (RFC 6266).
    from unicodedata import normalize, combining
    from urllib.parse import quote

    ascii_fallback = "".join(
        c for c in normalize('NFKD', filename) if not combining(c)
    ).encode('ascii', 'ignore').decode('ascii')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = (
        f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{quote(filename)}"
    )
    wb.save(response)
    return response


@admin.action(description="📊 Exporter la sélection en Excel (.xlsx)")
def export_registrations_excel(modeladmin, request, queryset):
    wb = _build_registrations_workbook(queryset)
    return _xlsx_response(wb, "inscriptions_evenements.xlsx")


# ─────────────────────────────
# Publication
# ─────────────────────────────
@admin.register(Publication)
class PublicationAdmin(admin.ModelAdmin):
    list_display   = ['title', 'author', 'category', 'publication_type', 'is_published', 'is_featured', 'views_count', 'registrations_link', 'created_at']
    list_filter    = ['is_published', 'is_featured', 'publication_type', 'category']
    list_editable  = ['is_published', 'is_featured']
    search_fields  = ['title', 'content']
    prepopulated_fields = {'slug': ('title',)}
    inlines        = [PublicationImageInline, PublicationVideoInline]
    exclude        = ['author']

    # ── Utilise CKEditor si disponible, sinon Textarea large ──
    def formfield_for_dbfield(self, db_field, request, **kwargs):
        if db_field.name == 'content':
            try:
                from ckeditor.widgets import CKEditorWidget
                kwargs['widget'] = CKEditorWidget(config_name='default')
            except ImportError:
                from django.forms import Textarea
                kwargs['widget'] = Textarea(attrs={'rows': 30, 'style': 'width:100%;font-size:14px;'})
        return super().formfield_for_dbfield(db_field, request, **kwargs)

    def get_queryset(self, request):
        # Tout admin ayant accès à ce module voit toutes les publications
        # (et pas seulement les siennes) — Django gère déjà les permissions
        # d'accès au module via is_staff / les permissions du modèle.
        return super().get_queryset(request)

    def save_model(self, request, obj, form, change):
        # L'auteur n'est jamais choisi dans le formulaire : c'est toujours
        # la personne connectée qui publie.
        if not obj.pk:
            obj.author = request.user
        super().save_model(request, obj, form, change)

    def cover_preview(self, obj):
        if obj.cover_image:
            return format_html('<img src="{}" style="height:50px;border-radius:6px;">', obj.cover_image.url)
        return "—"
    cover_preview.short_description = "Couverture"

    # ── Lien "fichier Excel" propre à chaque publication événement ──
    def registrations_link(self, obj):
        if obj.publication_type != 'evenement':
            return "—"
        count = obj.registrations.count()
        if not count:
            return "Aucune inscription"
        url = reverse('admin:publications_publication_export_registrations', args=[obj.pk])
        return format_html('<a href="{}">📊 Excel ({} inscrit{})</a>', url, count, "s" if count > 1 else "")
    registrations_link.short_description = "Inscriptions"

    def get_urls(self):
        urls = super().get_urls()
        custom = [
            path(
                '<int:pk>/export-registrations/',
                self.admin_site.admin_view(self.export_registrations_view),
                name='publications_publication_export_registrations',
            ),
        ]
        return custom + urls

    def export_registrations_view(self, request, pk):
        publication = self.get_object(request, pk)
        wb = _build_registrations_workbook(
            EventRegistration.objects.filter(publication=publication),
            sheet_title=publication.title,
        )
        safe_title = "".join(c if c.isalnum() else "_" for c in publication.title)[:60]
        return _xlsx_response(wb, f"inscriptions_{safe_title}.xlsx")

    fieldsets = (
        ('Informations principales', {
            'fields': ('title', 'slug', 'category')
        }),
        ('Type de publication', {
            'fields': ('publication_type',),
            'description': (
                "« Article / Actualité » : un article classique. "
                "« Événement » : affiche les infos de date/lieu et permet aux visiteurs "
                "de s'inscrire via un formulaire intégré au site (les inscrits sont "
                "exportables en Excel, individuellement par publication, depuis la liste "
                "des publications ou l'onglet « Inscriptions aux événements »)."
            ),
        }),
        ('Détails de l\'événement (si type = Événement)', {
            'fields': ('event_date', 'event_location', 'registration_open', 'registration_deadline'),
            'classes': ('collapse',),
        }),
        ('Médias de couverture', {
            'fields': ('cover_image', 'cover_video'),
            'classes': ('collapse',),
        }),
        ('Contenu', {
            'fields': ('content',),
        }),
        ('Publication', {
            'fields': ('is_published', 'is_featured'),
        }),
    )


# ─────────────────────────────
# Inscriptions aux événements
# ─────────────────────────────
@admin.register(EventRegistration)
class EventRegistrationAdmin(admin.ModelAdmin):
    list_display    = ['first_name', 'last_name', 'publication', 'email', 'full_phone', 'country', 'city', 'created_at']
    list_filter     = ['publication']
    search_fields   = ['first_name', 'last_name', 'email', 'phone_number']
    readonly_fields = ['first_name', 'last_name', 'email', 'phone_code', 'phone_number', 'country', 'city', 'publication', 'created_at']
    actions         = [export_registrations_excel]

    def full_phone(self, obj):
        return obj.full_phone
    full_phone.short_description = "WhatsApp"
